from io import BytesIO
from docx import Document
from aiogram import Bot
from typing import Optional, List
from mammoth import convert_to_html
from bs4 import BeautifulSoup
import asyncio
import aiohttp
import hashlib
import subprocess
import traceback
import os
import re

from config.conf import logger
from dbs.conf_redis import get_redis_client
from dbs.mongo import Group, mongo
from dbs.yandex_cloud import Cloud

class vgpgk:

    _url = "https://vgpgk.ru/raspisanie/vgpgk-zameny-1-korpus.doc?v=2023011141816"
    _doc = os.path.abspath('./files/zameni.doc')
    _docx = os.path.abspath('./files/zameni.docx')
    
    # _doc = os.path.abspath(r'.\files\zameni.doc')
    # _docx = os.path.abspath(r'.\files\zameni.docx')
    
    __not_updated = True
    client = get_redis_client()
    cloud = Cloud()
    
    @classmethod
    def _calculate_sha256(cls, filepath: str) -> str:
        sha256_hash = hashlib.sha256()
        try:
            with open(filepath, "rb") as f:
                for chunk in iter(lambda: f.read(), b""):
                    sha256_hash.update(chunk)
            logger.debug(f'Хеш посчитан для {filepath}\n {sha256_hash.hexdigest()}')
            return sha256_hash.hexdigest()
        
        except Exception as e:
            logger.error(f'Проблемы с вычислением хеша: {e}')
            
        
    @classmethod
    async def _save_hash(cls, hash_val: str) -> None:
        try:
            await cls.client.set('zam_hash', hash_val)
            logger.debug(f'Хеш сохранен в редисе {hash_val}')
        except Exception as e:
            logger.error(f'Пролемы с сохранением хеша: {e}')
            
    
    @classmethod
    async def _load_hash(cls) -> Optional[str]:
        try:
            hash_val = await cls.client.get('zam_hash')
            logger.debug(f"Хеш прочитан {hash_val}")
            return hash_val
        except Exception as e:
            logger.error(f"Ошибка при загрузке хеша: {e}")
            return None

    @classmethod
    async def download_replace(cls) -> bool:
        old_hash = await cls._load_hash()
        try:
            async with aiohttp.ClientSession() as session:
                async with session.get(cls._url) as response:
                    
                    response.raise_for_status()
                    logger.info('Отправлен запрос на получение замен')
                    
                    with open(cls._doc, 'wb') as f:
                        async for chunk in response.content.iter_chunked(1024):
                            f.write(chunk)
                            
                    logger.info('получен новый doc файл')
            
            new_hash = cls._calculate_sha256(cls._doc)

            if new_hash is None:
                logger.error('Проблемы с рассчетом новго хеша')
                return False
            if new_hash != old_hash:
                await cls._save_hash(new_hash)
                cls.convert_doc_to_docx(cls._doc, cls._docx)
                await cls.cloud.upload(cls._docx)
                cls.__not_updated = False
                logger.info(f'Замены обновлены\n{new_hash}\n{old_hash}')
                return True
            elif new_hash == old_hash:
                 logger.info('Замены не были изменены')
            else:
                logger.warning('Проблемы с заменами')
            cls.__not_updated = True
        except Exception as e:
            logger.error(f"Ошибка при скачивании файла: {e}")
            return False

    @staticmethod
    def convert_doc_to_docx(doc_path: str, docx_path: str) -> None:
        try:
            if not os.path.exists(doc_path):
                logger.error(f"Файл не найден: {doc_path}")

            # soffice = r"C:\Program Files\LibreOffice\program\soffice.exe"

            command = [
                "soffice",
                "--headless",
                "--convert-to",
                "docx",
                "--outdir",
                os.path.dirname(docx_path),
                doc_path,
            ]
            subprocess.run(command, check=True, capture_output=True)

            if not os.path.exists(docx_path):
                logger.error(f"Не удалось преобразовать {doc_path} в {docx_path}")

            logger.info(f"Успешно преобразовано: {doc_path} -> {docx_path}")
        except Exception as e:
            logger.error("Ошибка при конвертации файла")
            raise
    
    
    @classmethod
    def convert_docx_to_html(cls):
        with open(cls._docx, "rb") as f:
            result = convert_to_html(f)
            html = result.value
        return html

    @classmethod
    async def parse_replace_from_html(cls, html_str: str) -> dict[str, str]:
        soup = BeautifulSoup(html_str, 'html.parser')
        table = soup.find('table')
        data = []
        
        for row in table.find_all('tr'):
            col_data = []
            for col in row.find_all('td'):
                text = '\n'.join(p.text for p in col.find_all('p'))
                if text:
                    col_data.append(text)
            data.append(col_data)

        group_name = re.compile(r'^[А-Я]{2}-[0-9]{3}')
        replace = dict()
        
        for row_i in range(len(data)):
            col = data[row_i]
            for col_i in range(len(col)):
                if re.match(group_name, data[row_i][col_i][0:6]):
                    replace[data[row_i][col_i]] = data[row_i+1][col_i]
                    await cls.client.set(data[row_i][col_i], data[row_i+1][col_i])
        
        return replace
    
    
    @classmethod
    async def get_replace(cls, group_name: str) -> Optional[List[str]]:
        find = False
        
        buf = await cls.cloud.download()
        document = Document(buf)
        if await cls.client.get(group_name) and cls.__not_updated:
            zam = await cls.client.get(group_name)
            if zam:
                logger.debug(f"замены для {group_name} взяты из кеша: {zam}")
                return [group_name, zam]
        for table in document.tables:
            for row_i in range(len(table.rows)):
                for col_i in range(len(table.columns)):
                    logger.debug(table.cell(row_i, col_i).text)
                    ans = table.cell(row_i, col_i).text
                    logger.debug(f"ТЕКСТ: {ans}\nСТРОКА:{row_i} СТОБЕЦ {col_i}")
                    print(f"ТЕКСТ: {ans}\nСТРОКА:{row_i} СТОБЕЦ {col_i}")
                    if ans.startswith(group_name):
                        group = table.cell(row_i, col_i).text
                        replace = table.cell(row_i+1, col_i).text
                        find = True
                    if find:
                        logger.info(f"для группы {group} найдены замены {replace}")
                        await cls.client.set(group_name, replace, 86400)
                        return [group, replace]
                    
        logger.debug(f"ЗАМЕНЫ ДЛЯ {group_name} НЕ НАЙДЕНЫ")  
        return None



async def sheduled_replace(bot: Bot, interval: int = 1800):

    while True:
        try:
            await asyncio.sleep(interval)
            if await vgpgk.download_replace():
                html_content = vgpgk.convert_docx_to_html()
                replaces = await vgpgk.parse_replace_from_html(html_content)
                groups = await mongo.get_all_groups()
                for group in groups:
                    group_replace = replaces.get(group.group_name)
                    for chat in group.chats:
                        nani = 'Замен нет и это точно'
                        await bot.send_message(chat_id=chat, text=f'{group.group_name}\n{group_replace if group_replace else nani}')
                logger.info("Рассылка отправлена")
            else:
                # html_content = vgpgk.convert_docx_to_html()
                # replaces = await vgpgk.parse_replace_from_html(html_content)
                # groups = await mongo.get_all_groups()
                # for group in groups:
                #     group_replace = replaces.get(group.group_name)
                #     for chat in group.chats:
                #         nani = 'Замен нет и это точно'
                #         await bot.send_message(chat_id=chat, text=f'{group.group_name}\n{group_replace if group_replace else nani}')
                logger.info('Замены не изменились, рассылка не отправлена')

        except Exception:
            logger.error(f"Ошибка при автоматическом обновлении замен {traceback.format_exc()}")
            

            
if __name__=="__main__":
    replace_data = vgpgk.get_replace("ИС-223")
    if replace_data:
        print(replace_data[0])
        print(replace_data[1])
        print(vgpgk._docx)
    else:
        print('pusto')
        